from __future__ import annotations

import io
import math
import os
import re
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from tempfile import mkdtemp
from typing import Iterable, List, Optional, Sequence, Set

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from fin_stmts_parse.src.core.contracts import (
    StatementExcelArtifact,
    StatementMarkdownDocument,
)
from fin_stmts_parse.src.markdown_table_extractor import MarkdownTableExtractor
from fin_stmts_parse.utils.logger import logger


_SHEET_NAME_ILLEGAL_CHARS = re.compile(r'[\[\]:*?/\\]')
_MARKDOWN_SEPARATOR_RE = re.compile(r"^\|[\s\-:|]+\|$")
_MARKDOWN_STRONG_RE = re.compile(r"(?<!\\)(\*\*\*|___|\*\*|__)(.+?)(?<!\\)\1")
_MARKDOWN_INLINE_CODE_RE = re.compile(r"`([^`]*)`")
_HTML_BREAK_RE = re.compile(r"<br\s*/?>", re.IGNORECASE)
_HTML_TAG_RE = re.compile(r"<[^>]+>")
_DATE_LIKE_HEADER_RE = re.compile(
    r"^(?:19|20)\d{2}(?:\s*[年./-]\s*\d{1,2})?(?:\s*[月./-]\s*\d{1,2})?\s*(?:日)?$"
)
_AMOUNT_HEADER_KEYWORDS = (
    "金额",
    "余额",
    "发生额",
    "累计",
    "本期",
    "上期",
    "本年",
    "上年",
    "本月",
    "上月",
    "期末",
    "期初",
    "年末",
    "年初",
    "收入",
    "成本",
    "费用",
    "利润",
    "净额",
)
_NUMERIC_TEXT_RE = re.compile(
    r"^[+-]?(?:\d+|\d{1,3}(?:[,，]\d{3})+)(?:\.\d+)?$"
)
_AMOUNT_DECIMAL_QUANT = Decimal("0.01")
_AMOUNT_NUMBER_FORMAT = "#,##0.00"

_TEXT_COLOR = "1F2937"
_MUTED_TEXT_COLOR = "64748B"
_TABLE_TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
_TABLE_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_METADATA_LABEL_FILL = PatternFill("solid", fgColor="EEF2F7")
_METADATA_VALUE_FILL = PatternFill("solid", fgColor="F8FAFC")
_ALT_ROW_FILL = PatternFill("solid", fgColor="F8FBFF")
_WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
_THIN_SIDE = Side(style="thin", color="D9E2EC")
_MEDIUM_SIDE = Side(style="medium", color="A0AEC0")
_TABLE_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)
_HEADER_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_MEDIUM_SIDE, bottom=_THIN_SIDE)


@dataclass(frozen=True)
class _ExcelCellValue:
    value: str
    bold: bool = False


class FinancialStatementExcelExportCore:
    def __init__(self) -> None:
        self.logger = logger
        self.table_extractor = MarkdownTableExtractor(llm_client=None)

    def export_markdown_documents_to_excel(
        self,
        documents: Sequence[StatementMarkdownDocument],
        *,
        file_name: str,
        output_dir: Optional[str] = None,
    ) -> StatementExcelArtifact:
        resolved_output_dir = self._ensure_output_dir(output_dir)
        normalized_file_name = self._normalize_excel_file_name(file_name)

        workbook = Workbook()
        workbook.remove(workbook.active)

        sheet_names: List[str] = []
        exported_sheets: List[dict] = []
        skipped_tables: List[dict] = []

        for document in documents:
            if not document.content.strip():
                skipped_tables.append(
                    {
                        "table_name": document.table_name,
                        "reason": "no_table_content",
                    }
                )
                continue

            table_blocks = self.table_extractor.extract_all_tables_from_markdown(document.content)
            parsed_blocks = self._parse_table_blocks(table_blocks)
            if not parsed_blocks:
                skipped_tables.append(
                    {
                        "table_name": document.table_name,
                        "reason": "no_table_content",
                    }
                )
                continue

            worksheet = workbook.create_sheet(
                title=self._build_sheet_name(document.table_name, sheet_names)
            )
            sheet_names.append(worksheet.title)
            exported_sheets.append(
                {
                    "table_name": document.table_name,
                    "date": document.date,
                    "is_consolidated": document.is_consolidated,
                    **({"page": document.page} if document.page is not None else {}),
                }
            )
            self._write_document_sheet(worksheet, document, parsed_blocks)

        if not sheet_names:
            return StatementExcelArtifact(
                source_file_name=self._pick_source_file_name(documents, normalized_file_name),
                source_file_path=self._pick_source_file_path(documents),
                skipped_tables=skipped_tables,
                summary=self._build_summary(
                    input_markdown_count=len(documents),
                    generated_sheet_count=0,
                    skipped_table_count=len(skipped_tables),
                    generated_workbook_count=0,
                ),
            )

        excel_file_path = self._build_unique_output_path(resolved_output_dir, normalized_file_name)
        workbook.save(excel_file_path)

        return StatementExcelArtifact(
            source_file_name=self._pick_source_file_name(documents, normalized_file_name),
            source_file_path=self._pick_source_file_path(documents),
            excel_file_name=os.path.basename(excel_file_path),
            excel_file_path=excel_file_path,
            sheet_names=sheet_names,
            exported_sheets=exported_sheets,
            skipped_tables=skipped_tables,
            summary=self._build_summary(
                input_markdown_count=len(documents),
                generated_sheet_count=len(sheet_names),
                skipped_table_count=len(skipped_tables),
                generated_workbook_count=1,
            ),
        )

    def _ensure_output_dir(self, output_dir: Optional[str]) -> str:
        if isinstance(output_dir, str) and output_dir.strip():
            resolved = os.path.abspath(output_dir.strip())
            os.makedirs(resolved, exist_ok=True)
            return resolved
        return mkdtemp(prefix="fin_stmts_excel_")

    @staticmethod
    def _normalize_excel_file_name(file_name: str) -> str:
        candidate = Path(file_name or "input").name
        stem = Path(candidate).stem or "input"
        return f"{stem}.xlsx"

    @staticmethod
    def _pick_source_file_name(
        documents: Sequence[StatementMarkdownDocument],
        fallback: str,
    ) -> str:
        for document in documents:
            if document.source_file_name:
                return document.source_file_name
        return fallback

    @staticmethod
    def _pick_source_file_path(documents: Sequence[StatementMarkdownDocument]) -> str:
        for document in documents:
            if document.source_file_path:
                return document.source_file_path
        return ""

    def _build_unique_output_path(self, output_dir: str, file_name: str) -> str:
        base_path = os.path.join(output_dir, file_name)
        if not os.path.exists(base_path):
            return base_path

        stem = Path(file_name).stem or "input"
        suffix = Path(file_name).suffix or ".xlsx"
        counter = 2
        while True:
            candidate = os.path.join(output_dir, f"{stem}_{counter}{suffix}")
            if not os.path.exists(candidate):
                return candidate
            counter += 1

    @staticmethod
    def _build_sheet_name(table_name: str, used_names: Sequence[str]) -> str:
        candidate = _SHEET_NAME_ILLEGAL_CHARS.sub("_", (table_name or "Sheet").strip()) or "Sheet"
        candidate = candidate[:31] or "Sheet"
        if candidate not in used_names:
            return candidate

        counter = 2
        while True:
            suffix = f"_{counter}"
            deduped = f"{candidate[: 31 - len(suffix)]}{suffix}"
            if deduped not in used_names:
                return deduped
            counter += 1

    def _parse_table_blocks(self, table_blocks: Iterable[dict]) -> List[dict]:
        parsed_blocks: List[dict] = []
        for block in table_blocks:
            block_type = block.get("type")
            content = block.get("content", "")
            if block_type == "markdown":
                rows = self._parse_markdown_table(content)
            elif block_type == "html":
                rows = self._parse_html_table(content)
            else:
                rows = []

            if rows:
                parsed_blocks.append({"type": block_type, "rows": rows})
        return parsed_blocks

    def _parse_markdown_table(self, markdown_table: str) -> List[List[_ExcelCellValue]]:
        rows: List[List[_ExcelCellValue]] = []
        for raw_line in markdown_table.splitlines():
            line = raw_line.strip()
            if not line.startswith("|"):
                continue
            if _MARKDOWN_SEPARATOR_RE.match(line):
                continue
            cells = [self._parse_export_cell(cell.strip()) for cell in line.split("|")[1:-1]]
            if cells:
                rows.append(cells)
        return rows

    def _parse_html_table(self, html_table: str) -> List[List[_ExcelCellValue]]:
        try:
            dataframes = pd.read_html(io.StringIO(html_table), keep_default_na=False)
        except ValueError:
            self.logger.warning("HTML 表格未解析出任何 DataFrame")
            return []

        rows: List[List[_ExcelCellValue]] = []
        for dataframe in dataframes:
            normalized = dataframe.fillna("")
            header_row = self._build_html_header_row(normalized.columns)
            if header_row:
                rows.append(header_row)
            for row in normalized.itertuples(index=False):
                rows.append([self._parse_export_cell(value) for value in row])
        return rows

    def _build_html_header_row(self, columns: pd.Index) -> List[_ExcelCellValue]:
        if self._is_default_html_header(columns):
            return []

        header_row = [
            _ExcelCellValue(self._stringify_html_header_value(column), bold=True)
            for column in columns
        ]
        if any(cell.value for cell in header_row):
            return header_row
        return []

    @staticmethod
    def _is_default_html_header(columns: pd.Index) -> bool:
        if isinstance(columns, pd.RangeIndex):
            return True
        return list(columns) == list(range(len(columns)))

    def _stringify_html_header_value(self, value: object) -> str:
        if isinstance(value, tuple):
            parts: List[str] = []
            for item in value:
                normalized = self._stringify_cell(item)
                if normalized:
                    parts.append(normalized)
            return " / ".join(parts)
        return self._stringify_cell(value)

    def _write_document_sheet(
        self,
        worksheet,
        document: StatementMarkdownDocument,
        parsed_blocks: Sequence[dict],
    ) -> None:
        metadata_rows = [
            ["table_name", document.table_name],
            ["date", document.date],
            ["unit", document.unit],
            ["is_consolidated", document.is_consolidated],
            ["source_file_name", document.source_file_name],
            ["source_file_path", document.source_file_path],
        ]
        for row in metadata_rows:
            worksheet.append(row)

        worksheet.append([])

        table_ranges = []
        for index, block in enumerate(parsed_blocks, start=1):
            block_rows = block["rows"]
            block_width = max((len(row) for row in block_rows), default=1)
            amount_columns = self._detect_amount_columns(block_rows[0] if block_rows else [])
            worksheet.append([f"Table {index}"])
            title_row = worksheet.max_row
            header_row = None
            last_row = title_row
            for row_offset, row in enumerate(block_rows):
                row_number = worksheet.max_row + 1
                self._write_parsed_row(
                    worksheet,
                    row_number,
                    row,
                    amount_columns if row_offset > 0 else set(),
                )
                if row_offset == 0:
                    header_row = row_number
                last_row = row_number
            table_ranges.append(
                {
                    "title_row": title_row,
                    "header_row": header_row,
                    "last_row": last_row,
                    "last_col": max(block_width, 2),
                }
            )
            if index < len(parsed_blocks):
                worksheet.append([])

        self._style_document_sheet(worksheet, table_ranges)

    def _write_parsed_row(
        self,
        worksheet,
        row_number: int,
        row: Sequence[_ExcelCellValue],
        amount_columns: Set[int],
    ) -> None:
        for column_number, parsed_cell in enumerate(row, start=1):
            cell_value = parsed_cell.value
            number_format = None
            if column_number in amount_columns:
                parsed_number = self._parse_amount_number(parsed_cell.value)
                if parsed_number is not None:
                    cell_value = parsed_number[0]
                    number_format = parsed_number[1]
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
                value=cell_value,
            )
            if number_format is not None:
                cell.number_format = number_format
            if parsed_cell.bold:
                cell.font = Font(bold=True, color=_TEXT_COLOR)

    def _style_document_sheet(self, worksheet, table_ranges: Sequence[dict]) -> None:
        worksheet.sheet_view.showGridLines = False
        if table_ranges and table_ranges[0]["header_row"] is not None:
            worksheet.freeze_panes = f"B{table_ranges[0]['header_row'] + 1}"

        self._style_metadata_rows(worksheet)
        for table_range in table_ranges:
            self._style_table_range(worksheet, table_range)
        self._apply_column_widths(worksheet)
        self._apply_row_heights(worksheet, table_ranges)

    def _style_metadata_rows(self, worksheet) -> None:
        for row_number in range(1, 7):
            label_cell = worksheet.cell(row=row_number, column=1)
            value_cell = worksheet.cell(row=row_number, column=2)

            label_cell.font = Font(bold=True, color=_MUTED_TEXT_COLOR, size=10)
            label_cell.fill = _METADATA_LABEL_FILL
            label_cell.alignment = Alignment(horizontal="left", vertical="center")
            label_cell.border = _TABLE_BORDER

            value_cell.font = Font(color=_TEXT_COLOR, size=10)
            value_cell.fill = _METADATA_VALUE_FILL
            value_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            value_cell.border = _TABLE_BORDER

    def _style_table_range(self, worksheet, table_range: dict) -> None:
        last_col = table_range["last_col"]
        title_row = table_range["title_row"]
        header_row = table_range["header_row"]
        last_row = table_range["last_row"]

        for column_number in range(1, last_col + 1):
            cell = worksheet.cell(row=title_row, column=column_number)
            cell.fill = _TABLE_TITLE_FILL
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = _HEADER_BORDER

        if header_row is None:
            return

        for column_number in range(1, last_col + 1):
            cell = worksheet.cell(row=header_row, column=column_number)
            cell.fill = _TABLE_HEADER_FILL
            cell.font = Font(bold=True, color=_TEXT_COLOR, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = _HEADER_BORDER

        for row_number in range(header_row + 1, last_row + 1):
            fill = _ALT_ROW_FILL if (row_number - header_row) % 2 == 0 else _WHITE_FILL
            for column_number in range(1, last_col + 1):
                cell = worksheet.cell(row=row_number, column=column_number)
                markdown_bold = bool(cell.font and cell.font.bold)
                cell.fill = fill
                cell.font = Font(bold=markdown_bold, color=_TEXT_COLOR, size=10)
                cell.alignment = self._data_alignment(cell.value, column_number == 1)
                cell.border = _TABLE_BORDER

    @staticmethod
    def _data_alignment(value: object, is_first_column: bool) -> Alignment:
        if is_first_column:
            horizontal = "left"
        elif FinancialStatementExcelExportCore._looks_numeric_text(value):
            horizontal = "right"
        else:
            horizontal = "center"
        return Alignment(horizontal=horizontal, vertical="center", wrap_text=True)

    @staticmethod
    def _looks_numeric_text(value: object) -> bool:
        if value in ("", None):
            return False
        text = str(value).strip()
        if not text:
            return False
        normalized = (
            text.replace(",", "")
            .replace("，", "")
            .replace("%", "")
            .replace("−", "-")
            .strip()
        )
        if normalized.startswith("(") and normalized.endswith(")"):
            normalized = f"-{normalized[1:-1]}"
        return bool(re.fullmatch(r"-?\d+(\.\d+)?", normalized))

    @staticmethod
    def _detect_amount_columns(header_row: Sequence[_ExcelCellValue]) -> Set[int]:
        amount_columns: Set[int] = set()
        for column_number, cell in enumerate(header_row, start=1):
            if column_number == 1:
                continue
            header = re.sub(r"\s+", "", cell.value)
            if not header:
                continue
            if any(keyword in header for keyword in _AMOUNT_HEADER_KEYWORDS):
                amount_columns.add(column_number)
                continue
            if _DATE_LIKE_HEADER_RE.fullmatch(header):
                amount_columns.add(column_number)
        return amount_columns

    @staticmethod
    def _parse_amount_number(value: object) -> Optional[tuple[Decimal, str]]:
        if value in ("", None):
            return None
        text = str(value).strip()
        if not text:
            return None

        parenthesized_negative = text.startswith("(") and text.endswith(")")
        if parenthesized_negative:
            text = text[1:-1].strip()
        text = text.replace("−", "-").replace(" ", "")
        if not _NUMERIC_TEXT_RE.fullmatch(text):
            return None

        normalized = text.replace(",", "").replace("，", "")
        if parenthesized_negative:
            normalized = f"-{normalized}"

        try:
            number = Decimal(normalized).quantize(
                _AMOUNT_DECIMAL_QUANT,
                rounding=ROUND_HALF_UP,
            )
        except InvalidOperation:
            return None

        return number, _AMOUNT_NUMBER_FORMAT

    def _apply_column_widths(self, worksheet) -> None:
        for column_number in range(1, worksheet.max_column + 1):
            max_content_width = 0
            for row_number in range(1, worksheet.max_row + 1):
                value = worksheet.cell(row=row_number, column=column_number).value
                for line in self._stringify_cell(value).splitlines() or [""]:
                    max_content_width = max(max_content_width, self._display_width(line))

            if column_number == 1:
                min_width, max_width = 20, 48
            else:
                min_width, max_width = 14, 36
            width = min(max(max_content_width + 2, min_width), max_width)
            worksheet.column_dimensions[get_column_letter(column_number)].width = width

    def _apply_row_heights(self, worksheet, table_ranges: Sequence[dict]) -> None:
        title_rows = {table_range["title_row"] for table_range in table_ranges}
        header_rows = {
            table_range["header_row"]
            for table_range in table_ranges
            if table_range["header_row"] is not None
        }

        for row_number in range(1, worksheet.max_row + 1):
            row_values = [cell.value for cell in worksheet[row_number]]
            if not any(value not in ("", None) for value in row_values):
                worksheet.row_dimensions[row_number].height = 7.5
                continue

            if row_number in title_rows:
                min_height = 24
            elif row_number in header_rows:
                min_height = 22
            elif row_number <= 6:
                min_height = 20
            else:
                min_height = 18

            estimated_lines = self._estimate_wrapped_line_count(worksheet, row_number)
            worksheet.row_dimensions[row_number].height = min(
                max(min_height, estimated_lines * 15 + 4),
                90,
            )

    def _estimate_wrapped_line_count(self, worksheet, row_number: int) -> int:
        max_lines = 1
        for cell in worksheet[row_number]:
            text = self._stringify_cell(cell.value)
            if not text:
                continue
            column_width = worksheet.column_dimensions[
                get_column_letter(cell.column)
            ].width or 12
            usable_width = max(int(column_width) - 2, 8)
            cell_lines = 0
            for line in text.splitlines() or [""]:
                display_width = self._display_width(line)
                cell_lines += max(1, math.ceil(display_width / usable_width))
            max_lines = max(max_lines, cell_lines)
        return max_lines

    @staticmethod
    def _parse_export_cell(value: object) -> _ExcelCellValue:
        text = FinancialStatementExcelExportCore._stringify_cell(value)
        if not text:
            return _ExcelCellValue("")

        normalized = _HTML_BREAK_RE.sub("\n", text)
        normalized = _HTML_TAG_RE.sub("", normalized)
        bold = bool(_MARKDOWN_STRONG_RE.search(normalized))
        normalized = _MARKDOWN_STRONG_RE.sub(lambda match: match.group(2), normalized)
        normalized = _MARKDOWN_INLINE_CODE_RE.sub(lambda match: match.group(1), normalized)
        normalized = normalized.replace(r"\*", "*").replace(r"\_", "_")
        normalized = "\n".join(
            re.sub(r"[ \t]+", " ", line).strip()
            for line in normalized.splitlines()
        ).strip()
        return _ExcelCellValue(normalized, bold=bold)

    @staticmethod
    def _display_width(text: str) -> int:
        return sum(
            2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1
            for character in text
        )

    @staticmethod
    def _build_summary(
        *,
        input_markdown_count: int,
        generated_sheet_count: int,
        skipped_table_count: int,
        generated_workbook_count: int,
    ) -> dict:
        return {
            "input_file_count": 1,
            "input_markdown_count": input_markdown_count,
            "generated_sheet_count": generated_sheet_count,
            "skipped_table_count": skipped_table_count,
            "generated_workbook_count": generated_workbook_count,
        }

    @staticmethod
    def _stringify_cell(value: object) -> str:
        if value in ("", None):
            return ""
        return str(value)


def export_markdown_documents_to_excel(
    documents: Sequence[StatementMarkdownDocument],
    *,
    file_name: str,
    output_dir: Optional[str] = None,
) -> StatementExcelArtifact:
    return FinancialStatementExcelExportCore().export_markdown_documents_to_excel(
        documents,
        file_name=file_name,
        output_dir=output_dir,
    )
