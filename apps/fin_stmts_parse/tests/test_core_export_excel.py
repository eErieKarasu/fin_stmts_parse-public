import os
import tempfile
import unittest

from openpyxl import load_workbook

from apps.fin_stmts_parse.src.core.contracts import StatementMarkdownDocument
from apps.fin_stmts_parse.src.core.export_excel import (
    FinancialStatementExcelExportCore,
)


class FinancialStatementExcelExportCoreTest(unittest.TestCase):
    def setUp(self) -> None:
        self.core = FinancialStatementExcelExportCore()

    def test_export_markdown_table_writes_metadata_and_cells(self) -> None:
        document = StatementMarkdownDocument(
            table_name="资产负债表",
            page=12,
            date="202412",
            unit="元",
            is_consolidated=False,
            content=(
                "# 资产负债表\n\n"
                "| 项目 | 期末余额 |\n"
                "| --- | --- |\n"
                "| 货币资金 | 100 |\n"
            ),
            source_file_name="demo.pdf",
            source_file_path="/tmp/demo.pdf",
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            artifact = self.core.export_markdown_documents_to_excel(
                [document],
                file_name="demo.pdf",
                output_dir=temp_dir,
            )

            self.assertTrue(os.path.exists(artifact.excel_file_path))
            self.assertEqual(artifact.sheet_names, ["资产负债表"])
            self.assertEqual(
                artifact.exported_sheets,
                [
                    {
                        "table_name": "资产负债表",
                        "page": 12,
                        "date": "202412",
                        "is_consolidated": False,
                    }
                ],
            )
            workbook = load_workbook(artifact.excel_file_path)
            sheet = workbook["资产负债表"]

            self.assertEqual(sheet["A1"].value, "table_name")
            self.assertEqual(sheet["B1"].value, "资产负债表")
            self.assertEqual(sheet["A8"].value, "Table 1")
            self.assertEqual(sheet["A9"].value, "项目")
            self.assertEqual(sheet["B9"].value, "期末余额")
            self.assertEqual(sheet["A10"].value, "货币资金")
            self.assertEqual(sheet["B10"].value, 100)
            self.assertEqual(sheet["B10"].data_type, "n")
            self.assertEqual(sheet["B10"].number_format, "#,##0.00")

    def test_export_html_table_writes_cells(self) -> None:
        document = StatementMarkdownDocument(
            table_name="利润表",
            date="202412",
            content=(
                "# 利润表\n\n"
                "<table>"
                "<tr><td>项目</td><td>金额</td></tr>"
                "<tr><td>营业总收入</td><td>1000</td></tr>"
                "</table>"
            ),
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            artifact = self.core.export_markdown_documents_to_excel(
                [document],
                file_name="profit.pdf",
                output_dir=temp_dir,
            )

            workbook = load_workbook(artifact.excel_file_path)
            sheet = workbook["利润表"]
            self.assertEqual(sheet["A9"].value, "项目")
            self.assertEqual(sheet["B9"].value, "金额")
            self.assertEqual(sheet["A10"].value, "营业总收入")
            self.assertEqual(sheet["B10"].value, 1000)
            self.assertEqual(sheet["B10"].data_type, "n")
            self.assertEqual(sheet["B10"].number_format, "#,##0.00")

    def test_export_html_table_with_th_writes_header_row(self) -> None:
        document = StatementMarkdownDocument(
            table_name="利润表",
            date="202412",
            content=(
                "# 利润表\n\n"
                "<table>"
                "<thead><tr><th>项目</th><th>金额</th></tr></thead>"
                "<tbody><tr><td>营业总收入</td><td>1000</td></tr></tbody>"
                "</table>"
            ),
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            artifact = self.core.export_markdown_documents_to_excel(
                [document],
                file_name="profit.pdf",
                output_dir=temp_dir,
            )

            workbook = load_workbook(artifact.excel_file_path)
            sheet = workbook["利润表"]
            self.assertEqual(sheet["A9"].value, "项目")
            self.assertEqual(sheet["B9"].value, "金额")
            self.assertEqual(sheet["A10"].value, "营业总收入")
            self.assertEqual(sheet["B10"].value, 1000)
            self.assertEqual(sheet["B10"].data_type, "n")
            self.assertEqual(sheet["B10"].number_format, "#,##0.00")

    def test_export_applies_layout_and_markdown_bold_styles(self) -> None:
        document = StatementMarkdownDocument(
            table_name="资产负债表",
            date="202412",
            content=(
                "| 项目 | 期末余额 |\n"
                "| --- | ---: |\n"
                "| **货币资金** | **100** |\n"
                "| 一年内到期的非流动资产<br>其中：长期应收款 | 200 |\n"
            ),
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            artifact = self.core.export_markdown_documents_to_excel(
                [document],
                file_name="demo.pdf",
                output_dir=temp_dir,
            )

            workbook = load_workbook(artifact.excel_file_path)
            sheet = workbook["资产负债表"]

            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual(sheet.freeze_panes, "B10")
            self.assertGreaterEqual(sheet.column_dimensions["A"].width, 20)
            self.assertGreaterEqual(sheet.row_dimensions[11].height, 30)

            self.assertEqual(sheet["A9"].value, "项目")
            self.assertTrue(sheet["A9"].font.bold)
            self.assertEqual(sheet["A10"].value, "货币资金")
            self.assertTrue(sheet["A10"].font.bold)
            self.assertEqual(sheet["B10"].value, 100)
            self.assertEqual(sheet["B10"].data_type, "n")
            self.assertEqual(sheet["B10"].number_format, "#,##0.00")
            self.assertTrue(sheet["B10"].font.bold)
            self.assertEqual(sheet["A11"].value, "一年内到期的非流动资产\n其中：长期应收款")

    def test_export_converts_only_amount_columns_to_numeric_cells(self) -> None:
        document = StatementMarkdownDocument(
            table_name="资产负债表",
            date="202412",
            content=(
                "| 项目 | 行次 | 2024年12月31日 | 备注 |\n"
                "| --- | --- | ---: | --- |\n"
                "| 货币资金 | 1 | 1,234.567 | 1000 |\n"
            ),
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            artifact = self.core.export_markdown_documents_to_excel(
                [document],
                file_name="demo.pdf",
                output_dir=temp_dir,
            )

            workbook = load_workbook(artifact.excel_file_path)
            sheet = workbook["资产负债表"]

            self.assertEqual(sheet["B10"].value, "1")
            self.assertEqual(sheet["B10"].data_type, "s")
            self.assertEqual(sheet["C10"].value, 1234.57)
            self.assertEqual(sheet["C10"].data_type, "n")
            self.assertEqual(sheet["C10"].number_format, "#,##0.00")
            self.assertEqual(sheet["D10"].value, "1000")
            self.assertEqual(sheet["D10"].data_type, "s")

    def test_export_sanitizes_and_deduplicates_sheet_names(self) -> None:
        long_invalid_name = "利润/表*超长名称超长名称超长名称超长名称超长名称"
        documents = [
            StatementMarkdownDocument(
                table_name=long_invalid_name,
                date="202412",
                content="| 项目 | 金额 |\n| --- | --- |\n| A | 1 |",
            ),
            StatementMarkdownDocument(
                table_name=long_invalid_name,
                date="202409",
                is_consolidated=True,
                content="| 项目 | 金额 |\n| --- | --- |\n| B | 2 |",
            ),
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            artifact = self.core.export_markdown_documents_to_excel(
                documents,
                file_name="duplicated.pdf",
                output_dir=temp_dir,
            )

            self.assertEqual(len(artifact.sheet_names), 2)
            self.assertNotEqual(artifact.sheet_names[0], artifact.sheet_names[1])
            self.assertTrue(all(len(name) <= 31 for name in artifact.sheet_names))
            self.assertTrue(all("/" not in name and "*" not in name for name in artifact.sheet_names))
            self.assertEqual(
                artifact.exported_sheets,
                [
                    {
                        "table_name": long_invalid_name,
                        "date": "202412",
                        "is_consolidated": False,
                    },
                    {
                        "table_name": long_invalid_name,
                        "date": "202409",
                        "is_consolidated": True,
                    },
                ],
            )


if __name__ == "__main__":
    unittest.main()
