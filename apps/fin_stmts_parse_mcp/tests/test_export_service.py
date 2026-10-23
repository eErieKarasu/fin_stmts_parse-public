import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from lwade_utils.bootload import load_entry_module


os.environ["RUNNING_ENV"] = "local"

module = load_entry_module("apps/fin_stmts_parse_mcp", "mcp", "service.py")


class FinStmtsParseMarkdownToExcelTests(unittest.TestCase):
    def tearDown(self) -> None:
        module.export_service = None

    def test_export_generates_workbook_from_parse_results(self):
        parse_results = [
            {
                "file_name": "demo.pdf",
                "company_name": "测试公司",
                "markdown_files": [
                    {
                        "table_name": "资产负债表",
                        "page": 12,
                        "date": "202412",
                        "is_consolidated": False,
                        "content": (
                            "# 资产负债表\n\n"
                            "| 项目 | 期末余额 |\n"
                            "| --- | --- |\n"
                            "| 货币资金 | 100 |\n"
                        ),
                    }
                ],
            }
        ]
        uploaded_local_paths = []
        uploaded_object_keys = []

        def fake_upload(local_file_path, s3_file_path, config_key="oss", **kwargs):
            uploaded_local_paths.append(local_file_path)
            uploaded_object_keys.append(s3_file_path)
            self.assertTrue(
                s3_file_path.startswith(
                    "agent-financial_report_agent/fin_parse_excel/"
                )
            )
            self.assertEqual(config_key, "agent_oss")
            return "http://example.com/demo.xlsx"

        with tempfile.TemporaryDirectory() as temp_dir, patch.object(
            module.S3Client,
            "upload_file",
            side_effect=fake_upload,
        ):
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=parse_results,
                session_id="session-1",
                output_dir=temp_dir,
            )

            self.assertEqual(len(result), 1)
            item = result[0]
            self.assertEqual(
                item,
                {
                    "message": "",
                    "file_name": "demo.pdf",
                    "company_name": "测试公司",
                    "excel_file_name": "demo.xlsx",
                    "excel_file_id": uploaded_object_keys[0],
                    "excel_file_url": "http://example.com/demo.xlsx",
                    "sheets": [
                        {
                            "table_name": "资产负债表",
                            "page": 12,
                            "date": "202412",
                            "is_consolidated": "false",
                        }
                    ],
                },
            )
            self.assertEqual(len(uploaded_local_paths), 1)
            self.assertTrue(os.path.exists(uploaded_local_paths[0]))
            workbook = load_workbook(uploaded_local_paths[0])
            self.assertEqual(workbook.sheetnames, ["资产负债表"])

    def test_export_creates_multiple_sheets_for_one_input_file(self):
        parse_results = [
            {
                "file_name": "demo.pdf",
                "markdown_files": [
                    {
                        "table_name": "资产负债表",
                        "page": 12,
                        "date": "202412",
                        "is_consolidated": False,
                        "content": "| 项目 | 期末余额 |\n| --- | --- |\n| 货币资金 | 100 |",
                    },
                    {
                        "table_name": "利润表",
                        "page": 13,
                        "date": "202409",
                        "is_consolidated": True,
                        "content": "| 项目 | 金额 |\n| --- | --- |\n| 营业总收入 | 1000 |",
                    },
                ],
            }
        ]
        uploaded_local_paths = []
        uploaded_object_keys = []

        with tempfile.TemporaryDirectory() as temp_dir, patch.object(
            module.S3Client,
            "upload_file",
            side_effect=lambda local_file_path, s3_file_path, config_key="oss", **kwargs: (
                uploaded_local_paths.append(local_file_path)
                or uploaded_object_keys.append(s3_file_path)
                or "http://example.com/demo.xlsx"
            ),
        ):
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=parse_results,
                session_id="session-1",
                output_dir=temp_dir,
            )

            workbook = load_workbook(uploaded_local_paths[0])
            self.assertEqual(workbook.sheetnames, ["资产负债表", "利润表"])
            self.assertEqual(result[0]["excel_file_id"], uploaded_object_keys[0])
            self.assertEqual(
                result[0]["sheets"],
                [
                    {
                        "table_name": "资产负债表",
                        "page": 12,
                        "date": "202412",
                        "is_consolidated": "false",
                    },
                    {
                        "table_name": "利润表",
                        "page": 13,
                        "date": "202409",
                        "is_consolidated": "true",
                    },
                ],
            )

    def test_export_accepts_new_parse_output_contract(self):
        parse_results = [
            {
                "message": "解析成功",
                "file_name": "demo.pdf",
                "company_name": "公司名",
                "markdown_files": [
                    {
                        "table_name": "资产负债表",
                        "page": 12,
                        "date": "202412",
                        "is_consolidated": False,
                        "content": "| 项目 | 期末余额 |\n| --- | --- |\n| 货币资金 | 100 |",
                    }
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir, patch.object(
            module.S3Client,
            "upload_file",
            return_value="http://example.com/demo.xlsx",
        ):
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=parse_results,
                session_id="session-1",
                output_dir=temp_dir,
            )

            self.assertEqual(
                result[0],
                {
                    "message": "",
                    "file_name": "demo.pdf",
                    "company_name": "公司名",
                    "excel_file_name": "demo.xlsx",
                    "excel_file_id": result[0]["excel_file_id"],
                    "excel_file_url": "http://example.com/demo.xlsx",
                    "sheets": [
                        {
                            "table_name": "资产负债表",
                            "page": 12,
                            "date": "202412",
                            "is_consolidated": "false",
                        }
                    ],
                },
            )
            self.assertTrue(
                result[0]["excel_file_id"].startswith(
                    "agent-financial_report_agent/fin_parse_excel/"
                )
            )

    def test_export_infers_company_name_with_llm(self):
        parse_results = [
            {
                "file_name": "demo.pdf",
                "company_name": "",
                "markdown_files": [
                    {
                        "table_name": "资产负债表",
                        "date": "202412",
                        "is_consolidated": False,
                        "content": (
                            "# 资产负债表\n\n"
                            "**编制单位：** 测试公司\n\n"
                            "| 项目 | 期末余额 |\n"
                            "| --- | --- |\n"
                            "| 货币资金 | 100 |\n"
                        ),
                    }
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir, patch.object(
            module.CompanyNameResolver,
            "_extract_with_llm",
            return_value="测试公司",
        ), patch.object(
            module.S3Client,
            "upload_file",
            return_value="http://example.com/demo.xlsx",
        ):
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=parse_results,
                session_id="session-1",
                output_dir=temp_dir,
            )

        self.assertEqual(result[0]["company_name"], "测试公司")

    def test_export_falls_back_to_rules_when_llm_returns_empty(self):
        parse_results = [
            {
                "file_name": "demo.pdf",
                "company_name": "",
                "markdown_files": [
                    {
                        "table_name": "资产负债表",
                        "date": "202412",
                        "is_consolidated": False,
                        "content": (
                            "# 资产负债表\n\n"
                            "**编制单位：** 测试公司\n\n"
                            "| 项目 | 期末余额 |\n"
                            "| --- | --- |\n"
                            "| 货币资金 | 100 |\n"
                        ),
                    }
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir, patch.object(
            module.CompanyNameResolver,
            "_extract_with_llm",
            return_value="",
        ), patch.object(
            module.S3Client,
            "upload_file",
            return_value="http://example.com/demo.xlsx",
        ):
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=parse_results,
                session_id="session-1",
                output_dir=temp_dir,
            )

        self.assertEqual(result[0]["company_name"], "测试公司")

    def test_export_skips_when_no_markdown_files_exist(self):
        result = module.fin_stmts_parse_markdown_to_excel(
            parse_results=[
                {
                    "file_name": "empty.pdf",
                    "company_name": "空公司",
                    "markdown_files": [],
                }
            ],
            session_id="session-1",
        )

        self.assertEqual(
            result,
            [
                {
                    "message": "没有可导出的 markdown_files",
                    "file_name": "empty.pdf",
                    "company_name": "空公司",
                    "excel_file_name": "",
                    "excel_file_id": "",
                    "excel_file_url": "",
                    "sheets": [],
                }
            ],
        )

    def test_export_returns_message_when_no_exportable_tables_exist(self):
        with patch.object(module.S3Client, "upload_file") as mock_upload:
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=[
                    {
                        "file_name": "invalid.pdf",
                        "company_name": "公司名",
                        "markdown_files": [
                            {
                                "table_name": "资产负债表",
                                "date": "202412",
                                "is_consolidated": False,
                                "content": "这不是表格",
                            }
                        ],
                    }
                ],
                session_id="session-1",
            )

        mock_upload.assert_not_called()
        self.assertEqual(
            result,
            [
                {
                    "message": "没有可导出的财报表格",
                    "file_name": "invalid.pdf",
                    "company_name": "公司名",
                    "excel_file_name": "",
                    "excel_file_id": "",
                    "excel_file_url": "",
                    "sheets": [],
                }
            ],
        )

    def test_export_reports_upload_failure(self):
        parse_results = [
            {
                "file_name": "demo.pdf",
                "company_name": "公司名",
                "markdown_files": [
                    {
                        "table_name": "资产负债表",
                        "page": 12,
                        "date": "202412",
                        "is_consolidated": False,
                        "content": "| 项目 | 期末余额 |\n| --- | --- |\n| 货币资金 | 100 |",
                    }
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir, patch.object(
            module.S3Client,
            "upload_file",
            return_value=None,
        ):
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=parse_results,
                session_id="session-1",
                output_dir=temp_dir,
            )

        self.assertEqual(
            result,
            [
                {
                    "message": "Excel 导出成功但上传失败",
                    "file_name": "demo.pdf",
                    "company_name": "公司名",
                    "excel_file_name": "demo.xlsx",
                    "excel_file_id": "",
                    "excel_file_url": "",
                    "sheets": [
                        {
                            "table_name": "资产负债表",
                            "page": 12,
                            "date": "202412",
                            "is_consolidated": "false",
                        }
                    ],
                }
            ],
        )

    def test_export_defaults_to_daily_persistent_output_dir(self):
        service = module.FinStmtsExcelExportService.__new__(
            module.FinStmtsExcelExportService
        )
        service.export_core = Mock()
        service.company_name_resolver = Mock()
        service._upload_excel_file = Mock(
            return_value=("excel-id", "http://example.com/demo.xlsx")
        )
        service.export_core.export_markdown_documents_to_excel.return_value = Mock(
            excel_file_path="/tmp/demo.xlsx",
            excel_file_name="demo.xlsx",
            exported_sheets=[
                {
                    "table_name": "资产负债表",
                    "date": "202412",
                    "is_consolidated": False,
                }
            ],
        )

        with tempfile.TemporaryDirectory() as temp_dir, patch.dict(
            os.environ,
            {"MCP_PERSISTENT_BASE_DIR": temp_dir},
        ):
            service.fin_stmts_parse_markdown_to_excel(
                parse_results=[
                    {
                        "file_name": "demo.pdf",
                        "company_name": "公司名",
                        "markdown_files": [
                            {
                                "table_name": "资产负债表",
                                "date": "202412",
                                "is_consolidated": False,
                                "content": "| 项目 | 期末余额 |\n| --- | --- |\n| 货币资金 | 100 |",
                            }
                        ],
                    }
                ],
                session_id="session-1",
            )

            date_prefix = module.datetime.now().strftime("%Y%m%d")
            expected_output_dir = str(Path(temp_dir) / date_prefix / "output")
            call_kwargs = (
                service.export_core.export_markdown_documents_to_excel.call_args.kwargs
            )

        self.assertEqual(call_kwargs["output_dir"], expected_output_dir)

    def test_export_requires_session_id(self):
        with self.assertRaisesRegex(ValueError, "session_id 不能为空"):
            module.fin_stmts_parse_markdown_to_excel(
                parse_results=[
                    {
                        "file_name": "demo.pdf",
                        "company_name": "公司名",
                        "markdown_files": [],
                    }
                ],
                session_id="",
            )

    def test_tool_wrapper_forwards_parse_results_and_output_dir(self):
        expected = [
            {
                "message": "",
                "file_name": "demo.pdf",
                "company_name": "公司名",
                "excel_file_name": "demo.xlsx",
                "excel_file_id": "llm/fin_parse_excel/20260414/demo.xlsx",
                "excel_file_url": "http://example.com/demo.xlsx",
                "sheets": [
                    {
                        "table_name": "资产负债表",
                        "date": "202412",
                        "is_consolidated": "false",
                    }
                ],
            }
        ]
        mock_service = Mock()
        mock_service.fin_stmts_parse_markdown_to_excel.return_value = expected

        with patch.object(module, "get_export_service", return_value=mock_service):
            result = module.fin_stmts_parse_markdown_to_excel(
                parse_results=[{"file_name": "demo.pdf", "company_name": "公司名", "markdown_files": []}],
                session_id="session-1",
                output_dir="/tmp/out",
            )

        mock_service.fin_stmts_parse_markdown_to_excel.assert_called_once_with(
            parse_results=[{"file_name": "demo.pdf", "company_name": "公司名", "markdown_files": []}],
            session_id="session-1",
            output_dir="/tmp/out",
        )
        self.assertEqual(result, expected)


if __name__ == "__main__":
    unittest.main()
